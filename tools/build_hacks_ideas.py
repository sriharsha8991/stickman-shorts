#!/usr/bin/env python3
"""Builds the "Life & productivity hacks" idea bank (plan/hacks_ideas.csv + an .xlsx with two tabs).

Every row was checked against primary sources on 2026-10-21 (verdict + evidence strength + caveat).
Rows are IDEAS, not the production queue: promote one by copying it into content_plan.csv.
Usage: python3 tools/build_hacks_ideas.py [out.xlsx]
"""
import csv
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)

COLS = [
    "ID", "Status", "Category", "Format", "Title", "Hook overlay (frame 1)", "Hook line (caption)",
    "Mechanism / angle", "Visual metaphor", "Twist", "Payoff / takeaway", "Loop / bookend", "Length (s)",
    "Primary goal", "Verified finding", "Verdict", "Evidence", "Caveat (keep on screen / in script)",
    "Sources", "Risk flag", "Music mood",
]

HEALTH = "Health-edu (add 'Education, not medical advice'; no cures/claims of treatment)"
NONE = "-"
M_CUR = "Curious -> build -> reveal: lo-fi keys + marimba, 110-120 BPM"
M_PLAY = "Playful comedic bounce: pizzicato/marimba, 120-130 BPM"
M_RISE = "Rising pulse + riser into a big hit on the number, 120 BPM"
M_CALM = "Calm morning: soft piano + warm pad + light shaker, 80-90 BPM"
M_DOC = "Storybook documentary: piano + strings, 90-100 BPM"
M_FIGHT = "Fight energy: drums + brass stabs, 130 BPM"

R = []


def add(cat, fmt, title, overlay, hook, mech, meta, twist, payoff, loop, length, goal, finding, verdict, ev, caveat, src, risk, music):
    R.append([f"H{len(R) + 1:02d}", "Idea", cat, fmt, title, overlay, hook, mech, meta, twist, payoff, loop, length, goal,
              finding, verdict, ev, caveat, src, risk, music])


# ---------------------------------------------------------------- MOVE (get active)
MOVE = "Move · get active without a gym"
add(MOVE, "Myth vs Reality", "10,000 steps came from a pedometer, not a study", "10,000 STEPS? SAYS WHO?",
    "Your 10,000-step goal likely came from a 1960s pedometer's name, not a study.",
    "The number likely comes from the 1965 Japanese pedometer 'Manpo-kei' (10,000-steps meter); big studies show benefits level off far lower",
    "A giant pedometer box with a price tag reading MANPO-KEI 1965", "The real curve flattens around 7,000",
    "Aim for ~7,000; any extra steps still count", "Step counter ticks 9,999 -> 10,000 -> resets", 35, "Shares + Saves",
    "Lee 2019 JAMA IM: origin 'likely' the 1965 Manpo-kei pedometer. Ding 2025 Lancet Public Health (57 studies): 7,000 vs 2,000 steps/day linked to 47% lower all-cause mortality, 38% lower dementia",
    "MYTH (that 10k is science-based)", "Strong", "Say 'likely came from'; mortality figures are associations",
    "https://whish.stanford.edu/wp-content/uploads/2019/10/Step-Volume-Intensity-Mortality_JAMA-IntMedicine-2019-.pdf ; https://www.sydney.edu.au/news-opinion/news/2025/07/24/rethink-the-10000-a-day-step-goal-study-suggests.html (doi:10.1016/S2468-2667(25)00164-1)",
    HEALTH, M_CUR)
add(MOVE, "Number Reveal", "Your step sweet spot depends on your age", "HOW MANY STEPS IS ENOUGH?",
    "If you're under 60 and over 60, your 'enough' step count is different.",
    "Dose-response: risk drops steeply from the lowest step counts, then plateaus (6-8k for 60+, 8-10k for under 60)",
    "Two staircases that flatten into plateaus at different heights", "The biggest win is the first few thousand steps",
    "Add 1,000 steps a week until you hit your plateau", "Counter rolls back to the first step", 30, "Saves",
    "Paluch 2022 Lancet Public Health: 15 cohorts, 47,471 adults; highest vs lowest quartile HR 0.47; plateau ~6,000-8,000 (60+) and ~8,000-10,000 (<60)",
    "SUPPORTED", "Strong", "Observational; 'linked to', not 'causes'",
    "https://discovery.ucl.ac.uk/id/eprint/10170017/1/1-s2.0-S2468266721003029-main.pdf", HEALTH, M_RISE)
add(MOVE, "Mechanism Reveal", "One-minute huffing bursts", "3 MINUTES. NO GYM.",
    "Three one-minute bursts of huffing and puffing a day, hidden in your normal life.",
    "Vigorous intermittent lifestyle activity (VILPA): short daily bursts (stairs, fast walking, carrying shopping) push the heart rate up",
    "Spark turns the stairs, a bus sprint and grocery bags into three tiny 'boss levels'",
    "Non-exercisers who did ~3 bursts/day had far lower risk", "Turn 3 everyday moments into bursts: stairs, power-walk, heavy bags",
    "Third burst ends where the first began", 30, "Saves",
    "Stamatakis 2022 Nature Medicine: 25,241 UK Biobank non-exercisers; median 3 bouts/day (~4.4 min); ~3 bouts/day linked to 38-40% lower all-cause & cancer mortality, 48-49% lower CVD mortality",
    "SUPPORTED", "Moderate", "Single observational cohort; healthier people may move more",
    "https://www.ukbiobank.ac.uk/publications/association-of-wearable-device-measured-vigorous-intermittent-lifestyle-physical-activity-with-mortality/ (doi:10.1038/s41591-022-02100-x)",
    HEALTH, M_RISE)
add(MOVE, "Versus Fight", "Walk vs stand after lunch", "WALK vs STAND",
    "Standing desk or a short walk: which actually fights the after-lunch sugar spike?",
    "Muscle contraction pulls glucose out of the blood; light walking breaks beat standing breaks",
    "Boxing ring: WALK vs STAND, a sugar-spike meter between them", "Walking wins by more than double",
    "Every 20-30 min after a meal: 2-5 minutes of light walking", "Meter dips back to calm", 35, "Saves",
    "Buffey 2022 Sports Medicine meta-analysis (7 studies, n=461): vs sitting, light walking lowered post-meal glucose (ES -0.72), standing only -0.31; Dunstan 2012 Diabetes Care RCT (n=19): 2-min light walks every 20 min cut glucose 24%",
    "SUPPORTED", "Strong", "Repeated breaks over hours, short-term markers; not for managing diabetes",
    "https://researchrepository.ul.ie/server/api/core/bitstreams/da179fcd-eb47-4011-8127-5e735ba8ff4c/content (doi:10.1007/s40279-022-01649-4) ; doi:10.2337/dc11-1931",
    HEALTH, M_FIGHT)
add(MOVE, "Myth vs Reality", "Your standing desk isn't the fix", "STANDING = HEALTHY?",
    "You bought a standing desk to get healthier. The evidence says: not so fast.",
    "Sit-stand desks reduce sitting time, but health benefits are unproven; movement, not standing still, is what helps",
    "A stick figure frozen like a statue at a desk while a walking figure zooms past", "Standing still is still still",
    "Use the desk as a cue to walk, not to stand like a statue", "Desk rises, figure walks off", 30, "Shares",
    "Cochrane review (Shrestha 2016; 20 studies, 2,174 people): sit-stand desks cut sitting by 30 min-2 h per workday; evidence low to very low certainty, health effects 'very uncertain'",
    "MYTH (that it makes you healthier)", "Moderate", "Review is strong; the underlying studies are weak",
    "https://www.cochrane.org/about-us/news/health-effects-sit-stand-desks-and-interventions-aimed-reduce-sitting-work-are-still-unproven",
    HEALTH, M_PLAY)
add(MOVE, "Mechanism Reveal", "Take the stairs: 5 flights a day", "50 STAIRS A DAY",
    "Five flights of stairs a day was linked to a 20% lower heart-disease risk.",
    "Stairs are a free vertical workout hidden in every building",
    "A building where every staircase is a glowing heart meter", "You don't need a gym, you need a staircase",
    "Skip the lift for anything under 3 floors", "Elevator 'ding' and he takes the stairs again", 30, "Saves",
    "Song 2023 Atherosclerosis (UK Biobank, >450,000 adults, 12.5 y): 50+ stair steps/day linked to 20% lower risk of atherosclerotic CVD",
    "SUPPORTED", "Moderate", "Observational; 2024 '39%' figure is a conference abstract only",
    "doi:10.1016/j.atherosclerosis.2023.117300 ; https://www.technologynetworks.com/diagnostics/news/climbing-more-than-five-flights-of-stairs-a-day-reduces-heart-disease-risk-379360",
    HEALTH, M_CUR)
add(MOVE, "Number Reveal", "Weekend warrior still counts", "ONLY WORKOUT ON WEEKENDS?",
    "If you only exercise on weekends, you're not wasting your time.",
    "Total weekly activity matters more than how it's spread out",
    "A calendar where Mon-Fri are empty and Sat-Sun are on fire", "Weekend warriors had ~30% lower death risk vs inactive",
    "Hit 150 minutes any way you can, even in 2 sessions", "Calendar flips back to Monday", 30, "Shares",
    "O'Donovan 2017 JAMA IM (63,591 adults 40+): weekend warriors vs inactive, all-cause HR 0.70, CVD HR 0.60; WHO: 150-300 min/week moderate activity, 'every move counts'",
    "SUPPORTED", "Moderate", "Self-reported activity; association",
    "https://www.lboro.ac.uk/media-centre/press-releases/2017/january/weekend-warriors-fight-off-death-for-longer.html ; https://www.paho.org/en/news/26-11-2020-every-move-counts-towards-better-health-says-who",
    HEALTH, M_FIGHT)
add(MOVE, "Mechanism Reveal", "Gym-only audiobooks (temptation bundling)", "ONLY AT THE GYM",
    "Lock your favourite show so you can only play it while you work out.",
    "Temptation bundling: pair a 'want' with a 'should' so the reward pulls you to the habit",
    "A treasure chest (the guilty-pleasure show) chained to a treadmill", "The boost fades unless you keep the lock",
    "Pick one guilty pleasure and only allow it during exercise", "Chest locks again", 30, "Saves",
    "Milkman, Minson & Volpp 2014 Management Science field experiment: audiobooks only at the gym -> +51% gym visits initially; effect faded over time; 61% later paid for gym-only access",
    "PARTLY", "Moderate", "Initial boost that fades; university sample",
    "https://repository.upenn.edu/oid_papers/150 (doi:10.1287/mnsc.2013.1784)", NONE, M_PLAY)
add(MOVE, "Mechanism Reveal", "Walk to unlock ideas", "STUCK? STAND UP.",
    "Stuck on an idea? Your brain brainstorms better when your legs are moving.",
    "Walking boosts divergent (brainstorming) thinking, not single-answer problem solving",
    "A lightbulb that only glows while the treadmill belt moves", "It helps brainstorming, not maths",
    "Take brainstorm calls as walking calls", "Bulb flickers back on", 30, "Shares",
    "Oppezzo & Schwartz 2014 JEP:LMC, 4 experiments, 176 participants: 81% improved on creative (alternate uses) task while walking; convergent thinking did not improve",
    "PARTLY", "Weak-Moderate", "Small lab studies; don't say '60%' without context",
    "https://www.apa.org/news/press/releases/2014/04/creativity-walk (doi:10.1037/a0036577)", NONE, M_CUR)
add(MOVE, "Number Reveal", "120 minutes of nature a week", "2 HOURS. ONE WEEK.",
    "There's a weekly dose of nature where people start feeling better: 120 minutes.",
    "Threshold effect: people with 120+ min/week in nature report better health and wellbeing",
    "A weekly bar slowly filling with leaves until it clicks at 120", "Under 120 showed no link; it can be split into short visits",
    "Book 2 hours of green time across the week", "Bar empties for a new week", 30, "Saves",
    "White 2019 Scientific Reports (n=19,806, England): 120+ min/week linked to good health (OR 1.59) and high wellbeing (OR 1.23); no association below 120; peak 200-300 min",
    "SUPPORTED", "Moderate", "Cross-sectional, self-reported; a threshold seen in data, not a proven dose",
    "doi:10.1038/s41598-019-44097-3 ; https://www.scimex.org/newsfeed/two-hours-a-week-in-nature-can-make-you-happier-and-healthier",
    HEALTH, M_CALM)
add(MOVE, "Mechanism Reveal", "Fidgeting burns more than you think (NEAT)", "WHY SOME PEOPLE DON'T GAIN",
    "Same extra food, very different weight gain. The difference was tiny movements.",
    "Non-exercise activity thermogenesis: fidgeting, pacing, standing burns energy without 'exercise'",
    "Two figures eat the same cake; one keeps tapping, pacing, fidgeting", "Some people automatically move more when they overeat",
    "Pace on calls, stand to talk, take the long way", "Foot taps the beat that started the video", 30, "Shares",
    "Levine, Eberhardt & Jensen 1999 Science: 16 adults overfed 1,000 kcal/day for 8 weeks; NEAT rose by up to 692 kcal/day and explained differences in fat gain",
    "PARTLY", "Weak", "n=16; doesn't prove deliberate fidgeting makes you lose weight",
    "https://pubmed.ncbi.nlm.nih.gov/9880251/ ; https://sciencedaily.com/releases/1999/01/990119080449.htm",
    HEALTH, M_PLAY)

# ---------------------------------------------------------------- FOCUS (productivity)
FOCUS = "Focus · productivity"
add(FOCUS, "Mechanism Reveal", "The if-then trick", "IF ___ THEN ___",
    "One sentence that makes you way more likely to actually do the thing.",
    "Implementation intentions hand the decision to a cue in advance, so the situation triggers the action automatically",
    "A domino: the cue (IF) knocks over the action (THEN) without effort", "You're not adding willpower, you're removing a decision",
    "Write one: 'If it's 7 AM and I've made coffee, then I walk 10 minutes'", "Domino resets", 30, "Saves",
    "Gollwitzer & Sheeran 2006 meta-analysis (94 tests): d = 0.65 on goal attainment; Milne 2002 RCT (n=248): 91% exercised with a when/where plan vs 29-39% without (secondary report)",
    "SUPPORTED", "Strong", "Mostly short-term, self-reported; don't say 'doubles your success'",
    "https://kops.uni-konstanz.de/entities/publication/2e749bfb-8533-437c-8203-7e788c910c5f ; doi:10.1348/135910702169420",
    NONE, M_CUR)
add(FOCUS, "Myth vs Reality", "21 days to form a habit is a myth", "21 DAYS?",
    "If you quit your new habit on day 22, you were never on a 21-day clock.",
    "Habits become automatic gradually; average ~66 days, huge spread; missing one day barely matters",
    "A progress bar labelled 21 that keeps stretching into a long road", "The range was 18 to 254 days",
    "Stop counting to 21; miss a day, don't miss two", "Day counter 1 again", 35, "Shares + Saves",
    "Lally 2010 Eur J Soc Psych: 82 completers, 84 days; average ~66 days to automaticity, range 18-254 (upper end extrapolated); a single missed day had little effect",
    "MYTH (21 days)", "Moderate", "Small sample; simple behaviours",
    "https://bps.org.uk/research-digest/how-form-habit", NONE, M_PLAY)
add(FOCUS, "Myth vs Reality", "Willpower isn't a battery", "WILLPOWER: 3% LEFT",
    "You've been told willpower runs out like a phone battery. A 23-lab test said: not so fast.",
    "The famous 'ego depletion' effect failed a big preregistered replication",
    "A battery icon draining... then the battery turns out to be a sticker", "d = 0.04 across 2,141 people",
    "Make hard tasks easier to start instead of 'saving' willpower", "Battery shows 100% again", 35, "Shares",
    "Hagger et al. 2016 Perspectives on Psych Science: preregistered 23-lab replication, N=2,141, effect d = 0.04 [-0.07, 0.15]",
    "MYTH", "Strong", "Tested one paradigm; debate continues",
    "https://www.psychologicalscience.org/journals/perspectives/1745691616652873/", NONE, M_CUR)
add(FOCUS, "Number Reveal", "The Monday effect is real (fresh starts)", "+33% ON MONDAYS",
    "There's a reason your new plan always starts on a Monday.",
    "Temporal landmarks feel like a clean slate, so people start goals after them",
    "A calendar where every Monday, birthday and new month is a glowing door", "New semester was the biggest: +47%",
    "Use any landmark (a Monday, the 1st) as your launch day, then plan for the dip", "Next door glows", 30, "Saves",
    "Dai, Milkman & Riis 2014 Management Science: 11,912 gym members; visits up after a new week (+33.4%), new month (+14.4%), new year (+11.6%), new semester (+47.1%), birthday (+7.5%)",
    "SUPPORTED", "Moderate", "Shows when people start, not whether they keep going",
    "https://faculty.wharton.upenn.edu/wp-content/uploads/2014/06/Dai_Fresh_Start_2014_Mgmt_Sci.pdf", NONE, M_RISE)
add(FOCUS, "Myth vs Reality", "The '23 minutes to refocus' stat", "23 MINUTES TO REFOCUS?",
    "That '23 minutes to refocus after a ping' stat? It's not what the research said.",
    "The 23 min came from an interview about time until returning to a task; the study found interrupted people work FASTER but more stressed",
    "A stopwatch that splits into a stress meter", "Interruptions don't slow you, they stress you",
    "Batch notifications; protect one deep-work block", "Ping! loop", 30, "Shares",
    "Gloria Mark, Gallup interview 2006: interrupted work resumed in ~23 min 15 s (time to return); Mark CHI 2008 (n=48): interrupted people finished faster with no quality loss but more stress and frustration",
    "MYTH (as usually worded)", "Moderate", "Keep the nuance: the number is real but means something else",
    "https://news.gallup.com/businessjournal/23146/Too-Many-Interruptions-Work.aspx ; https://www.interruptions.net/literature/Mark-CHI08.pdf",
    NONE, M_PLAY)
add(FOCUS, "Mechanism Reveal", "Attention residue: why switching is so tiring", "YOUR BRAIN IS STILL IN THE LAST TAB",
    "Part of your brain is still in the last thing you were doing.",
    "Unfinished tasks leave 'attention residue' that drags on the next task; switching has a real cost",
    "Sticky threads trailing from one task box to the next", "Even finishing isn't enough unless you close it cleanly",
    "Before switching: write the next step down, then close the tab", "Threads snap clean", 30, "Saves",
    "Leroy 2009 OBHDP (2 experiments): leaving a task unfinished hurt performance on the next; Rubinstein, Meyer & Evans 2001: measurable task-switching costs; '40%' is Meyer's spoken estimate (APA)",
    "SUPPORTED", "Moderate", "Say 'researchers estimate up to 40%' or drop the number",
    "https://ideas.repec.org/a/eee/jobhdp/v109y2009i2p168-181.html ; https://www.apa.org/topics/research/multitasking",
    NONE, M_CUR)
add(FOCUS, "Mechanism Reveal", "Unfinished tasks nag you, until you plan them", "WHY CAN'T I STOP THINKING ABOUT IT?",
    "That task you didn't finish keeps popping into your head. There's a one-line fix.",
    "Unfulfilled goals produce intrusive thoughts; making a specific plan quiets them (Zeigarnik-style effect)",
    "Pop-up notifications inside the head that close when a plan card is written", "You don't have to finish it, just plan it",
    "Write the next concrete step + when", "Last pop-up closes", 30, "Saves",
    "Masicampo & Baumeister 2011 JPSP (6 studies): unfulfilled goals caused intrusive thoughts and worse task performance; making a specific plan eliminated the effects",
    "SUPPORTED", "Weak-Moderate", "One lab, no direct replication found",
    "https://supp.apa.org/psycarticles/supplemental/a0024192/a0024192_supp.html", NONE, M_CUR)
add(FOCUS, "Mechanism Reveal", "To-do list before bed = faster sleep", "CAN'T SLEEP? WRITE THIS.",
    "Five minutes of writing tomorrow's to-dos helped people fall asleep faster.",
    "Offloading tomorrow's tasks onto paper reduces bedtime rumination",
    "Thoughts buzzing like flies get caught in a notebook", "Writing what's NOT done beat writing what IS done",
    "Write a specific to-do list before bed", "Notebook closes, lights out", 30, "Saves",
    "Scullin 2018 J Exp Psych: General (n=57, polysomnography): to-do list group fell asleep ~9 min faster than completed-tasks group; more specific lists -> faster",
    "SUPPORTED", "Weak-Moderate", "One small study of young adults, one night",
    "https://news.web.baylor.edu/news/story/2018/can-writing-your-dos-help-you-doze-baylor-study-suggests-jotting-down-tasks-can",
    HEALTH, M_CALM)
add(FOCUS, "Myth vs Reality", "Your phone isn't stealing your IQ (mostly)", "YOUR PHONE STEALS YOUR BRAIN?",
    "The viral 'your phone drains your brain even when it's off' study. Here's what happened next.",
    "A famous 2017 finding mostly failed to replicate; a small effect on working memory may remain",
    "A phone with a vacuum hose to the head... then the hose is mostly air", "The famous version is overstated",
    "Still worth it: phone in another room removes the urge to check", "Phone slides off the desk", 35, "Shares",
    "Ward 2017 JACR: phone in another room -> better capacity scores; Hartmann 2021 (N=302) no effect; Parry 2024 meta-analysis (k=56, N=7,093): only small working-memory effect",
    "PARTLY", "Moderate", "Don't claim IQ loss; frame as small and debated",
    "https://www.bps.org.uk/research-digest/revisiting-brain-drain-effect ; https://forrt.org/flora-replication-atlas/doi/10.1086/691462/",
    NONE, M_PLAY)
add(FOCUS, "Mechanism Reveal", "Micro-breaks: recharge, don't expect miracles", "10 MINUTES OFF",
    "A break under 10 minutes really does recharge you. Just don't expect it to make you smarter.",
    "Micro-breaks raise vigour and cut fatigue; performance gains are small and mostly on easy tasks",
    "A battery charging during a mini coffee break", "Feel better: yes. Work better on hard tasks: not reliably",
    "Schedule breaks before you're exhausted", "Battery icon loops", 30, "Saves",
    "Albulescu 2022 PLOS ONE meta-analysis (22 samples, N=2,335): vigour d=.36, fatigue reduction d=.35, performance d=.16 (n.s.)",
    "PARTLY", "Moderate", "Benefits are mostly about energy and mood",
    "https://doaj.org/article/c6f37b6ba159469296b2cb968b65505f", NONE, M_CALM)
add(FOCUS, "Versus Fight", "Pomodoro vs 'I'll break when I'm tired'", "TIMER vs VIBES",
    "Scheduled breaks vs 'I'll stop when I'm tired'. One of them leaves you more drained.",
    "Fixed-interval breaks reduced fatigue and distraction compared with self-chosen breaks",
    "Boxing ring: a tomato timer vs a sleepy cloud", "Same amount of work got done, but the timer group felt better",
    "Try 25/5 (or 50/10) blocks for a study day", "Timer rings again", 30, "Saves",
    "Biwer et al. 2023 Brit J Educ Psych (n=87): Pomodoro-style breaks -> less fatigue/distraction, better concentration and motivation than self-regulated breaks; task completion did not differ",
    "PARTLY", "Weak", "One small study; mood/focus, not output",
    "https://cris.maastrichtuniversity.nl/en/publications/understanding-effort-regulation-comparing-pomodoro-breaks-and-sel",
    NONE, M_FIGHT)
add(FOCUS, "Number Reveal", "The pilot power nap", "26 MINUTES",
    "NASA let long-haul pilots nap in the cockpit. Here's what about 26 minutes of sleep did.",
    "A short planned nap reduces attention lapses and micro-sleeps",
    "A cockpit where the pilot dozes while a co-pilot keeps watch, alertness meter rising", "The famous '34% / 54%' numbers aren't in the report",
    "Keep naps short (10-30 min) and early afternoon", "Alarm beeps, eyes open", 30, "Shares",
    "Rosekind et al. NASA 1994/95 (21 pilots): nappers slept ~26 min; reaction-time lapses 81 vs 124; EEG micro-events 2.90 vs 6.37 per pilot",
    "PARTLY", "Weak-Moderate", "Tiny sample of pilots; don't quote 34%/54%",
    "https://ntrs.nasa.gov/citations/19950006379", HEALTH, M_RISE)
add(FOCUS, "Mechanism Reveal", "Set your own deadlines (evenly spaced)", "DUE: SOMEDAY",
    "'I'll do it by the end of the month' is the most dangerous deadline you can set.",
    "People use self-imposed deadlines to beat procrastination, but they space them badly; evenly spaced deadlines worked best",
    "A cliff edge at month-end vs a staircase of small deadlines", "Self-set deadlines help, but less than evenly spaced ones",
    "Break the big deadline into weekly mini-deadlines", "Next step lights up", 30, "Saves",
    "Ariely & Wertenbroch 2002 Psychological Science: self-imposed costly deadlines reduced procrastination but performed worse than evenly spaced external deadlines",
    "PARTLY", "Moderate", "Abstract-level numbers only",
    "https://scholars.duke.edu/publication/810365", NONE, M_CUR)
add(FOCUS, "POV Meme", "POV: the 2-minute rule", "TAKES 2 MINUTES?",
    "POV: you spent 20 minutes planning a task that takes 2 minutes.",
    "Practitioner heuristic (David Allen, Getting Things Done): if it takes under 2 minutes, do it now",
    "A tiny task growing into a monster in the to-do list", "Planning it took longer than doing it",
    "Under 2 minutes? Do it now", "The list shrinks to zero", 18, "Shares",
    "Heuristic from David Allen's 'Getting Things Done' (2001); no experimental test found",
    "HEURISTIC", "None (rule of thumb)", "Present as a rule of thumb, not science",
    "https://www.shortform.com/blog/the-two-minute-rule/ (secondary)", NONE, M_PLAY)

# ---------------------------------------------------------------- RESET (sleep, energy, mood)
RESET = "Reset · sleep, energy & mood"
add(RESET, "Mechanism Reveal", "Your 3 PM coffee is stealing an hour of sleep", "3 PM COFFEE",
    "Coffee 6 hours before bed cut sleep by over an hour, and people didn't notice.",
    "Caffeine blocks adenosine (the sleep-pressure signal); average half-life ~4 h (2-8 h), so afternoon cups are still active at night",
    "A coffee cup's ghost following the figure into bed and holding the eyelids open", "They didn't feel the lost sleep",
    "Set a caffeine cut-off time (e.g., 8+ hours before bed)", "Clock ticks back to 3 PM", 35, "Shares + Saves",
    "Drake 2013 J Clin Sleep Med (n=12): 400 mg caffeine 6 h before bed reduced total sleep by >1 h; EFSA: average half-life ~4 h (2-8 h)",
    "SUPPORTED", "Moderate", "400 mg dose; people metabolise caffeine very differently",
    "https://aasm.org/late-afternoon-and-early-evening-caffeine-can-disrupt-sleep-at-night/ ; https://www.efsa.europa.eu/sites/default/files/corporate_publications/files/efsaexplainscaffeine150527.pdf",
    HEALTH, M_CUR)
add(RESET, "Myth vs Reality", "You can't catch up on sleep at the weekend", "SLEEP DEBT: PAID?",
    "You think Saturday's lie-in pays back the week's sleep debt. Your body disagrees.",
    "Weekend recovery sleep didn't undo the metabolic effects of weekday sleep loss",
    "A piggy bank labelled SLEEP that won't accept weekend coins", "Catch-up group did worse on insulin sensitivity",
    "Protect a consistent sleep window on weekdays", "Alarm on Monday again", 35, "Shares",
    "Depner 2019 Current Biology (n=36, 2-week lab study): insulin sensitivity fell 13% with steady 5-h nights and 27% with 5-h nights plus weekend catch-up",
    "MYTH", "Moderate", "Small lab study; metabolic markers over 2 weeks",
    "https://www.nih.gov/news-events/nih-research-matters/weekend-catch-cant-counter-chronic-sleep-deprivation",
    HEALTH, M_CUR)
add(RESET, "Number Reveal", "5 minutes of sighing", "5 MINUTES. JUST BREATHE OUT.",
    "Five minutes of a breathing pattern with long exhales beat meditation for mood in one study.",
    "Cyclic sighing (double inhale, long exhale) slows breathing and calms the body",
    "Two balloons filling (double inhale), then a long slow deflate", "It beat mindfulness on daily mood in a small trial",
    "Try 5 minutes: two short inhales, one long exhale", "Balloon fills again", 30, "Saves",
    "Balban 2023 Cell Reports Medicine RCT (n=108, 28 days): cyclic sighing raised daily positive affect more than mindfulness meditation (1.89 vs 1.22)",
    "SUPPORTED (preliminary)", "Weak-Moderate", "Small, not preregistered; not a treatment for anxiety/depression",
    "https://pmc.ncbi.nlm.nih.gov/articles/PMC9873947", HEALTH, M_CALM)
add(RESET, "Mechanism Reveal", "Morning light sets your body clock", "OPEN THE CURTAINS",
    "Bright days and dark nights: your body clock is literally reading the light.",
    "Daytime light anchors the circadian rhythm; light at night pushes it later",
    "A sun-powered clock inside the head that runs fast or slow with light", "Night light was linked to worse mental health, day light to better",
    "Get outside in the morning; dim screens & lights before bed", "Sun rises again", 30, "Saves",
    "Burns 2023 Nature Mental Health (86,631 UK Biobank adults, wrist light sensors): more daytime light linked to lower risk of depression and other disorders; more night light to higher risk",
    "SUPPORTED", "Moderate", "Association; 'linked to', not 'cures'",
    "https://www.nationalelfservice.net/mental-health/depression/associations-between-daily-light-exposure-and-psychiatric-disorders/",
    HEALTH, M_CALM)
add(RESET, "Myth vs Reality", "8 glasses of water a day", "8 GLASSES. SAYS WHO?",
    "Nobody can find the study behind '8 glasses of water a day'.",
    "Total water comes from all drinks and food; thirst is an adequate guide for most healthy people",
    "Eight glasses lined up like soldiers... and a watermelon walks in", "About 20% of your water comes from food",
    "Drink to thirst; more in heat or exercise", "Glass refills", 30, "Shares",
    "Valtin 2002 Am J Physiol review: no scientific evidence for 8x8 in healthy adults; National Academies 2004: total water ~3.7 L men / 2.7 L women from all food and drink, ~20% from food",
    "MYTH", "Strong", "More needed in heat/exercise/pregnancy; never suggest drinking less",
    "https://geiselmed.dartmouth.edu/news/2002_h2/08aug2002_water.shtml ; https://www.nationalacademies.org/news/report-sets-dietary-intake-levels-for-water-salt-and-potassium-to-maintain-health-and-reduce-chronic-disease-risk",
    HEALTH, M_PLAY)
add(RESET, "Number Reveal", "Friends are a health habit", "50%",
    "Strong relationships were linked to 50% better odds of survival.",
    "Social connection is a health factor comparable to well-known risk factors",
    "A web of strings holding a figure up; cut strings make it wobble", "Social ties ranked alongside quitting smoking in the analysis",
    "Text one friend today; make one plan this week", "Phone buzzes: a friend", 30, "Shares + Saves",
    "Holt-Lunstad, Smith & Layton 2010 PLoS Medicine meta-analysis (148 studies, 308,849 people): stronger social relationships OR 1.50 for survival",
    "SUPPORTED", "Strong", "Observational; say '50% higher odds of survival', not 'live 50% longer'",
    "https://journals.plos.org/plosmedicine/article/info:doi/10.1371/journal.pmed.1000316", HEALTH, M_RISE)
add(RESET, "Mechanism Reveal", "Spend ₹ on someone else", "SPEND IT ON THEM",
    "Spending a small amount on someone else made people happier than spending it on themselves.",
    "Prosocial spending: giving creates a warm-glow mood boost",
    "Two shopping bags: one for me, one for you; the 'you' bag glows", "It was the small amount, not the big one",
    "Buy a friend a chai this week", "Chai cup returns", 30, "Shares",
    "Dunn, Aknin & Norton 2008 Science: spending a small amount on others raised happiness more than on self; Curry 2018 meta-analysis (27 studies, N=4,045): kindness δ = 0.28",
    "SUPPORTED", "Moderate", "Modest effect; keep money purely educational",
    "https://hbswk.hbs.edu/item/5944.html ; https://ora.ox.ac.uk/objects/uuid:4701365b-5760-4ece-a390-857a5f7b3c0c",
    "Finance-edu (education only; no products)", M_CUR)
add(RESET, "Myth vs Reality", "Gratitude journals: small boost, not magic", "GRATITUDE = HAPPINESS?",
    "Gratitude journaling won't change your life overnight. Here's what it actually does.",
    "Small mood benefit vs doing nothing; about the same as other positive activities",
    "A journal with a tiny sun that rises a little each entry", "It doesn't beat other active exercises",
    "Use it if you like it; 3 specific good things a night", "Pen taps the page", 30, "Saves",
    "Kirca 2023 meta-analysis (25 RCTs, ~6,745 people): g = 0.22 vs neutral controls, g = 0.12 (n.s.) vs active controls; Cregg & Cheavens 2021: small effect on depression/anxiety symptoms",
    "PARTLY", "Moderate", "Not a treatment for depression",
    "https://link.springer.com/10.1007/s41042-023-00086-6 ; https://ideas.repec.org/a/spr/jhappi/v22y2021i1d10.1007_s10902-020-00236-6.html",
    HEALTH, M_CALM)
add(RESET, "Myth vs Reality", "Cold showers: fewer sick days, not fewer colds", "COLD SHOWER CHALLENGE",
    "People who ended showers cold took 29% fewer sick days. But they weren't sick less.",
    "30-90 s of cold at the end of a shower; people felt more energetic and called in sick less, though illness days didn't change",
    "A thermometer that drops while a 'SICK DAY' calendar loses pages", "Same illness days, fewer days off",
    "Try the last 30 seconds cold (if you're healthy)", "Tap turns warm again", 30, "Shares",
    "Buijze 2016 PLOS ONE RCT (n=3,018): sickness absence down 29% (IRR 0.71); no difference in days of illness",
    "PARTLY", "Moderate", "Self-reported, unblinded; avoid with heart conditions",
    "https://doaj.org/article/4c0517724c1747f69283a44cc6c53e40", HEALTH, M_PLAY)
add(RESET, "Mechanism Reveal", "Screens in bed delay your body clock", "ONE MORE SCROLL",
    "Reading on a glowing screen before bed made people take longer to fall asleep.",
    "Bright light in the evening suppresses melatonin and shifts the body clock later",
    "Melatonin 'sleep fairies' scattering from a glowing tablet", "The effect was ~10 minutes, not catastrophic",
    "Switch to paper or dim, warm screens for the last hour", "Screen dims to black", 30, "Saves",
    "Chang et al. PNAS 2015 (n=12, lab crossover, 4 h/night e-reader vs print): ~10 min longer to fall asleep, less REM, suppressed melatonin, less alert next morning",
    "PARTLY", "Weak-Moderate", "Tiny lab study with long bright reading; avoid 'blue light ruins sleep'",
    "https://www.psu.edu/news/research/story/light-emitting-e-readers-detrimentally-shift-circadian-clock",
    HEALTH, M_CALM)
add(RESET, "Mechanism Reveal", "Heat wrecks deep sleep", "TOO HOT TO SLEEP?",
    "If you wake up a lot on hot nights, it's not just you. Heat steals deep sleep.",
    "The body needs to drop its core temperature to sleep; heat and humidity increase wakefulness and reduce deep and REM sleep",
    "A thermostat dial inside the figure's chest that must turn down to unlock sleep", "Cold bedrooms (with blankets) didn't hurt sleep stages",
    "Keep the bedroom cool and airy; light bedding", "Fan spins, loop", 30, "Saves",
    "Okamoto-Mizuno & Mizuno 2012 J Physiol Anthropol review: heat increases wakefulness and cuts slow-wave and REM sleep, humidity worsens it",
    "PARTLY", "Moderate", "Don't quote a precise ideal temperature as science",
    "https://jphysiolanthropol.biomedcentral.com/track/pdf/10.1186/1880-6805-31-14", HEALTH, M_CUR)
add(RESET, "Myth vs Reality", "Fake it till you make it (the smile study)", "SMILE = HAPPY?",
    "The famous 'hold a pen in your teeth to feel happier' study failed. Then the story got twistier.",
    "Facial feedback: posing a smile may slightly lift mood; the pen method didn't replicate, a 2022 global study found a small effect for posed smiles",
    "A pen in the teeth gets crossed out; a real posed smile wins", "Busted, then partly un-busted",
    "A genuine posed smile might help a little; don't expect magic", "Smile returns", 35, "Shares",
    "Wagenmakers 2016 registered replication (17 labs): pen effect ~0; Coles 2022 Many Smiles (3,878 people, 19 countries): voluntary posed smiles raised happiness, pen-in-mouth unreliable",
    "PARTLY", "Moderate", "Effect is small",
    "https://news.stanford.edu/stories/2022/10/posing-smiles-can-brighten-mood ; https://oasis.library.unlv.edu/psychology_fac_articles/507/",
    NONE, M_PLAY)


def main():
    out_csv = os.path.join(ROOT, "plan", "hacks_ideas.csv")
    with open(out_csv, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(COLS)
        w.writerows(R)
    print(f"wrote {out_csv} ({len(R)} ideas)")
    if len(sys.argv) > 1:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        wb = Workbook()
        ws = wb.active
        ws.title = "Hacks ideas"
        ws.append(COLS)
        for r in R:
            ws.append(r)
        hdr = PatternFill("solid", fgColor="0FB5A6")
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = hdr
            c.alignment = Alignment(wrap_text=True, vertical="top")
        widths = {"A": 6, "B": 8, "C": 22, "D": 16, "E": 34, "F": 26, "G": 44, "H": 44, "I": 36, "J": 32, "K": 34,
                  "L": 22, "M": 8, "N": 12, "O": 60, "P": 16, "Q": 12, "R": 36, "S": 60, "T": 30, "U": 30}
        for k, v in widths.items():
            ws.column_dimensions[k].width = v
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "F2"
        ws.auto_filter.ref = ws.dimensions
        g = wb.create_sheet("How to use")
        notes = [
            ["Life & productivity hacks: idea bank (researched 2026-10-21)"],
            [""],
            ["What this is", "Ideas for Spark videos about hacks that make people more active and feel better. Not the production queue: content_plan.csv still drives the runs."],
            ["Columns", "Same story columns as the content plan, plus Verified finding, Verdict, Evidence and Caveat from the research pass."],
            ["Verdict", "SUPPORTED = the claim holds · PARTLY = true with an important nuance · MYTH = the popular version is false (great for Myth vs Reality) · HEURISTIC = a rule of thumb, not research."],
            ["Evidence", "Strong = meta-analysis / large RCTs / official body · Moderate = large cohort or decent RCT · Weak = single small study."],
            ["Health rule", "Any row flagged Health-edu needs the on-screen line 'Education, not medical advice', 'linked to' wording for observational studies, and no promises of cures."],
            ["To produce one", "Reply to the production email with 'ADD H07'. The next run queues it (after any Redo rows, before the next plan row), re-verifies the facts and produces it as H07."],
            ["Counts", f"{len(R)} ideas: {sum(r[2] == MOVE for r in R)} Move, {sum(r[2] == FOCUS for r in R)} Focus, {sum(r[2] == RESET for r in R)} Reset."],
        ]
        for n in notes:
            g.append(n)
        g["A1"].font = Font(bold=True, size=14)
        for r in range(3, len(notes) + 1):
            g.cell(r, 1).font = Font(bold=True)
            g.cell(r, 2).alignment = Alignment(wrap_text=True, vertical="top")
        g.column_dimensions["A"].width = 18
        g.column_dimensions["B"].width = 110
        wb.save(sys.argv[1])
        print(f"wrote {sys.argv[1]}")


if __name__ == "__main__":
    main()
